# -*- coding: utf-8 -*-
"""PPT 第三轮：页脚错字（子任务务）与 P10 数据集体例统一。"""
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation

SRC = r'C:\Users\31908\Desktop\gpcr (1).pptx'
p = Presentation(SRC)
n_fix = 0

def fix_tf(tf, old, new, tag):
    global n_fix
    for para in tf.paragraphs:
        joined = "".join(r.text for r in para.runs)
        if old in joined:
            nb = joined.replace(old, new)
            if para.runs:
                para.runs[0].text = nb
                for r in para.runs[1:]:
                    r.text = ""
            n_fix += 1
            print("  %-10s %s -> %s" % (tag, old, new))

def walk(shapes, slide_no):
    for i, sh in enumerate(shapes):
        if sh.shape_type is not None and str(sh.shape_type).startswith("GROUP"):
            walk(sh.shapes, slide_no)
        elif sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    fix_tf(cell.text_frame, "子任务务", "子任务", "P%d tbl" % slide_no)
                    fix_tf(cell.text_frame, "B2AR 归档", "β2-AR 已归档", "P%d tbl" % slide_no)
        elif sh.has_text_frame:
            fix_tf(sh.text_frame, "子任务务", "子任务", "P%d[%d]" % (slide_no, i))
            fix_tf(sh.text_frame, "B2AR 归档", "β2-AR 已归档", "P%d[%d]" % (slide_no, i))

for n, slide in enumerate(p.slides, 1):
    walk(slide.shapes, n)
p.save(SRC)
print("\n共修正 %d 处；SAVED" % n_fix)
