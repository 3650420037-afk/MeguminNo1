# -*- coding: utf-8 -*-
"""P31 附录总表重建：删除被烘焙成矢量图的旧靶点标签，补写可见文本（标签 + 产出/入库数字）。

背景：该 PPT 由设计工具导出，部分文字是矢量轮廓（无法用文本替换），
      可编辑文本框的 run0 为 alpha=0（不可见），旧数字在后续 run 中。
"""
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = r'C:\Users\31908\Desktop\gpcr (1).pptx'
TEAL = RGBColor(0x33, 0x7A, 0x74)      # 与原矢量标签/数值一致
GRAY = RGBColor(0xE5, 0xE7, 0xEB)      # 原 产出/入库 数字色

p = Presentation(SRC)
s = p.slides[30]

# 1) 删除错误的矢量标签（行2 β2-AR、行3 D3、行4 5-HT2B）
victims = [(28, 3, "β2-AR"), (28, 2, "D3"), (38, 2, "5-HT2B")]
for gi, ci, what in victims:
    g = s.shapes[gi]
    kid = g.shapes[ci]
    kid._element.getparent().remove(kid._element)
    print("删除矢量标签: group%d kid%d (%s)" % (gi, ci, what))

# 2) 删除不可见的旧文本框（含错标签与旧数字）
for idx in sorted([24, 29, 34, 39, 47], reverse=True):
    sh = s.shapes[idx]
    print("删除不可见文本框 [%d] %r" % (idx, sh.text_frame.text.strip()[:30]))
    sh._element.getparent().remove(sh._element)


def add_text(x_cm, y_cm, w_cm, h_cm, txt, color, size, align=PP_ALIGN.LEFT, bold=False):
    tb = s.shapes.add_textbox(Cm(x_cm), Cm(y_cm), Cm(w_cm), Cm(h_cm))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = txt
    run.font.size = Pt(size)
    run.font.name = "Arial"
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


# 3) 行标签（行1 A2A 的矢量标签正确，保留）
add_text(1.61, 5.50, 2.4, 0.42, "D3", TEAL, 11)
add_text(1.61, 6.72, 2.4, 0.42, "5-HT2B", TEAL, 11)
add_text(1.61, 9.24, 2.0, 0.42, "合计", TEAL, 11)

# 4) 产出 / 入库 数字（列中心 5.58 / 8.04 cm）
rows = [("583", "395", 4.30), ("254", "106", 5.54), ("314", "206", 6.76), ("1,151", "707", 9.28)]
for out_v, in_v, y in rows:
    add_text(4.58, y, 2.0, 0.42, out_v, GRAY, 9.5, PP_ALIGN.CENTER)
    add_text(7.04, y, 2.0, 0.42, in_v, GRAY, 9.5, PP_ALIGN.CENTER)

p.save(SRC)
print("\nSAVED", SRC)
