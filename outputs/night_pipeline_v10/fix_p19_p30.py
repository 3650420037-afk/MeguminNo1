# -*- coding: utf-8 -*-
"""P19: Top5 表换 v10 数据（保留列对齐的空格 run）；排序口径说明改正。
   P30: 第2行量化证据缩短，避免折行压到下一行。"""
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.util import Cm

SRC = r'C:\Users\31908\Desktop\gpcr (1).pptx'
p = Presentation(SRC)

# ---------- P19 ----------
s = p.slides[18]
rows = {
    56: ("-12.88", "0.574", "2.8", "490", "1.295"),
    60: ("-12.77", "0.601", "4.2", "476", "1.177"),
    66: ("-12.53", "0.637", "3.3", "471", "1.306"),
    69: ("-12.15", "0.700", "2.7", "414", "1.431"),
    73: ("-12.10", "0.675", "2.5", "428", "1.425"),
}
for idx, vals in rows.items():
    sh = s.shapes[idx]
    para = sh.text_frame.paragraphs[0]
    vi = 0
    for r in para.runs:
        if r.text.strip():                      # 只改数值 run，空格 run 保留（列对齐靠 spc）
            r.text = vals[vi]
            vi += 1
    print("P19 [%d] -> %s (原 %r)" % (idx, " ".join(vals), sh.text_frame.text.strip()[:40]))

# 表头 "综合分" -> "类药性综合分"，并加宽
hdr = s.shapes[46]
hdr.width = Cm(2.3)
for para in hdr.text_frame.paragraphs:
    for r in para.runs:
        if r.text.strip():
            r.text = "类药性综合分"
print("P19 表头 ->", hdr.text_frame.text.strip(), "w=%.2fcm" % (hdr.width / 360000))

note = s.shapes[75]
for para in note.text_frame.paragraphs:
    for r in para.runs:
        if r.text.strip():
            r.text = "排序口径：对接分升序取 Top5；类药性综合分 = QED + 1 − SA/10；Top50 含 40 个唯一骨架"
print("P19 说明 ->", note.text_frame.text.strip())

# ---------- P30 ----------
s30 = p.slides[29]
tb = [sh for sh in s30.shapes if sh.has_table][0].table
cell = tb.cell(2, 2)
tf = cell.text_frame
for para in tf.paragraphs:
    for r in para.runs:
        if r.text.strip():
            r.text = "Top50最大相似 0.495"
print("P30 r2c2 ->", cell.text.strip())

p.save(SRC)
print("SAVED", SRC)
