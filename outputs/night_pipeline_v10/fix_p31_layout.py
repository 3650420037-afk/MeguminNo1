# -*- coding: utf-8 -*-
"""P31 收尾：加宽"对接均值"列避免折行；把合计行上移到空出的第 4 行位置，分隔线上移。"""
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.util import Cm, Emu

SRC = r'C:\Users\31908\Desktop\gpcr (1).pptx'
p = Presentation(SRC)
s = p.slides[30]

# 1) 加宽 "对接均值" 列文本框（x≈9.98，w 0.73/0.38 -> 1.5cm）
for sh in s.shapes:
    x, w = Emu(sh.left).cm, Emu(sh.width).cm
    if 9.9 <= x <= 10.2 and w < 0.8:
        sh.left, sh.width = Cm(9.95), Cm(1.5)
        print("加宽均值列 [%s] w=%.2f -> 1.50" % (sh.shape_id, w))

# 2) 合计行上移 1.26cm，分隔线上移到 7.90cm
for sh in s.shapes:
    y = Emu(sh.top).cm
    t = sh.text_frame.text.strip() if sh.has_text_frame else ""
    if 9.2 <= y <= 9.4 and (t.startswith(("合计", "1,151", "707", "-9.70", "-12.88", "42.0%"))):
        sh.top = Emu(sh.top - Cm(1.26))
        print("上移合计行 [%s] %r -> y=%.2f" % (sh.shape_id, t[:14], Emu(sh.top).cm))
    elif abs(y - 8.88) < 0.05 and sh.shape_type is not None and "FREEFORM" in str(sh.shape_type):
        sh.top = Cm(7.90)
        print("分隔线上移 [%s] -> y=7.90" % sh.shape_id)

# 3) 清掉第 4 行遗留的空文本框
for sh in list(s.shapes):
    if sh.has_text_frame and not sh.text_frame.text.strip() and 7.95 <= Emu(sh.top).cm <= 8.5 and Emu(sh.left).cm > 9.9:
        sh._element.getparent().remove(sh._element)
        print("删除空文本框 [%s]" % sh.shape_id)

p.save(SRC)
print("SAVED", SRC)
