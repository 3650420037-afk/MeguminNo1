# -*- coding: utf-8 -*-
"""P29「项目历程」原为半成品（只有标题与时间轴，无内容）——按真实项目历元补写时间线。"""
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SRC = r'C:\Users\31908\Desktop\gpcr (1).pptx'
TEAL = RGBColor(0x2F, 0xBF, 0xA8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY = RGBColor(0xC9, 0xD2, 0xDC)

MILES = [
    ("08-26", "基线打通", "Pocket2Mol 官方权重在 A2A 口袋跑通「采样 → 过滤 → 对接」闭环"),
    ("08-29", "官方权重库建成", "四靶点 8,904 个采样分子 → 5,864 个入库化合物"),
    ("09-21", "库级评估体系成形", "ROC 富集 5.80×、ADMET 分级、选择性 SI、Pareto 前沿与六维推荐"),
    ("09-27", "高密度数据集与架构改造", "12,244 对口袋–配体数据集落盘；终止模块扩容 4.2×、编码器解冻"),
    ("09-29", "v10 微调完成并晋级默认权重", "同种子成对十臂：4 更好 / 4 打平 / 2 更差，均值 −0.15 kcal/mol"),
    ("09-29", "引导束搜索定型", "λ=3、w_div=0.5；QED 中位 0.667 → 0.803（+20%），两种子复现"),
    ("09-30", "采样侧尺寸控制 + 药库重出", "三靶点 1,151 采样 → 707 入库；B2AR 生成入口按决策禁用"),
]

p = Presentation(SRC)
s = p.slides[28]
top0, gap = 2.35, 1.20

def box(x, y, w, h, txt, color, size, bold=False):
    tb = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    run = para.add_run(); run.text = txt
    run.font.size = Pt(size); run.font.bold = bold
    run.font.name = "Microsoft YaHei"
    run.font.color.rgb = color
    return tb

for i, (date, title, desc) in enumerate(MILES):
    y = top0 + i * gap
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Cm(1.10), Cm(y + 0.09), Cm(0.24), Cm(0.24))
    dot.fill.solid(); dot.fill.fore_color.rgb = TEAL
    dot.line.fill.background()
    dot.shadow.inherit = False
    box(1.55, y, 2.0, 0.40, date, TEAL, 10.5, True)
    box(3.35, y, 9.0, 0.40, title, WHITE, 12, True)
    box(1.55, y + 0.46, 11.0, 0.38, desc, GRAY, 9.5)
    print("P29 第 %d 条 @y=%.2f  %s %s" % (i + 1, y, date, title))

p.save(SRC)
print("SAVED")
