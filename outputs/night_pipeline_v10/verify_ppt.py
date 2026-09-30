# -*- coding: utf-8 -*-
"""PPT 全文校验：扫描残留旧数 / B2AR / 四靶点 / 口径错误。"""
import re, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation

SRC = r'C:\Users\31908\Desktop\gpcr (1).pptx'
p = Presentation(SRC)

BAD = [r"四靶点", r"四个\s*GPCR", r"β2", r"B2AR", r"b2ar", r"\b917\b", r"\b951\b", r"1,065",
       r"4,231", r"2,931", r"1,577", r"1,571", r"1,525", r"10,004", r"6,864",
       r"113\s*批", r"29\s*批", r"25\s*批", r"-9\.68", r"-13\.51", r"-13\.97", r"-8\.40",
       r"-11\.39", r"-9\.34", r"-12\.34", r"38\.8%", r"5\.5%", r"24\.9%", r"0\.187", r"0\.321",
       r"0\.32", r"<0\.56", r"98\.9%", r"56\s*$", r"POCKET2MOL", r"子任", r"化合物库"]
BADR = [(x, re.compile(x)) for x in BAD]

def texts(shapes, slide_no, out):
    for i, sh in enumerate(shapes):
        if sh.shape_type is not None and str(sh.shape_type).startswith("GROUP"):
            texts(sh.shapes, slide_no, out)
        elif sh.has_table:
            for r, row in enumerate(sh.table.rows):
                for c, cell in enumerate(row.cells):
                    if cell.text.strip():
                        out.append(("P%d tbl r%dc%d" % (slide_no, r, c), cell.text.strip()))
        elif sh.has_text_frame and sh.text_frame.text.strip():
            out.append(("P%d [%d]" % (slide_no, i), sh.text_frame.text.strip()))

hits = 0
for n, slide in enumerate(p.slides, 1):
    items = []
    texts(slide.shapes, n, items)
    for tag, t in items:
        flat = t.replace("\n", " / ")
        for pat, rx in BADR:
            if rx.search(flat):
                print("  !! P%-3d %-14s %-14s :: %s" % (n, tag.split(" ", 1)[-1], pat, flat[:110]))
                hits += 1
print("\n可疑残留命中: %d" % hits)
