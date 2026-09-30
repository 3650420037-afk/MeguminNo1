# -*- coding: utf-8 -*-
"""扫描全文：带 alpha=0（不可见）但含文字的 run —— 这些是"改了也看不见"的文本。"""
import re, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.util import Emu

p = Presentation(r'C:\Users\31908\Desktop\gpcr (1).pptx')
pat = re.compile(r'<a:srgbClr val="([0-9A-Fa-f]{6})">\s*<a:alpha val="(\d+)"/>')

def walk(shapes, n, path=""):
    for i, sh in enumerate(shapes):
        tag = "%s[%d]" % (path, i)
        if sh.shape_type is not None and str(sh.shape_type).startswith("GROUP"):
            walk(sh.shapes, n, tag + ".")
            continue
        if sh.has_table:
            for r, row in enumerate(sh.table.rows):
                for c, cell in enumerate(row.cells):
                    for para in cell.text_frame.paragraphs:
                        for run in para.runs:
                            if run.text.strip():
                                m = pat.search(run._r.xml)
                                if m and int(m.group(2)) == 0:
                                    print("P%-3d tbl r%dc%d INVISIBLE color=%s :: %r" % (n, r, c, m.group(1), run.text[:60]))
        elif sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if run.text.strip():
                        m = pat.search(run._r.xml)
                        if m and int(m.group(2)) == 0:
                            print("P%-3d %-10s y=%5.2f x=%5.2f INVISIBLE color=%s :: %r" % (
                                n, tag, Emu(sh.top).cm, Emu(sh.left).cm, m.group(1), run.text[:70]))
for n, slide in enumerate(p.slides, 1):
    walk(slide.shapes, n)
