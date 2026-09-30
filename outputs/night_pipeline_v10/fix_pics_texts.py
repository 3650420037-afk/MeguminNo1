# -*- coding: utf-8 -*-
"""替换 5 张旧数据光栅图 + 收尾文字/框宽修正。"""
import os, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.util import Cm, Emu

SRC = r'C:\Users\31908\Desktop\gpcr (1).pptx'
CH = os.path.join(r'D:\MMModel', 'outputs', 'night_pipeline_v10', 'charts')
p = Presentation(SRC)

def swap_picture(slide, idx, new_png, tag):
    sh = slide.shapes[idx]
    left, top, w, h = sh.left, sh.top, sh.width, sh.height
    pic = slide.shapes.add_picture(os.path.join(CH, new_png), left, top, width=w, height=h)
    sh._element.addnext(pic._element)
    sh._element.getparent().remove(sh._element)
    print("%s: %s 已替换 (%.2fx%.2fcm)" % (tag, new_png, Emu(w).cm, Emu(h).cm))

swap_picture(p.slides[14], 16, "p15_batch.png", "P15 批次统计图")
swap_picture(p.slides[15], 27, "p16_funnel.png", "P16 漏斗图")
swap_picture(p.slides[16], 10, "p17_qed.png", "P17 QED 分布图")
swap_picture(p.slides[19], 4, "p20_cluster.png", "P20 化学空间散点")
swap_picture(p.slides[19], 6, "p20_tanimoto.png", "P20 Tanimoto 分布")

def set_cell(tb, r, c, val, tag):
    cell = tb.cell(r, c)
    tf = cell.text_frame
    done = False
    for para in tf.paragraphs:
        for run in para.runs:
            if not done:
                run.text = val; done = True
            else:
                run.text = ""
    if not done:
        tf.text = val
    print("  %s r%dc%d -> %r" % (tag, r, c, val))

# ---------- P15 ----------
tb = [s for s in p.slides[14].shapes if s.has_table][0].table
set_cell(tb, 1, 4, "三靶点合计", "P15")
set_cell(tb, 2, 4, "1,151", "P15")
set_cell(tb, 3, 0, "11 批 × 50", "P15")
s15 = p.slides[14]
for idx, val in [(7, "5 批 × 50"), (8, "6 批 × 50"), (9, "22 批 × 50（实际 50-59/批）")]:
    sh = s15.shapes[idx]
    done = False
    for para in sh.text_frame.paragraphs:
        for run in para.runs:
            if not done:
                run.text = val; done = True
            else:
                run.text = ""
    print("  P15 [%d] -> %r" % (idx, val))

# ---------- P16 ----------
s16 = p.slides[15]
for para in s16.shapes[2].text_frame.paragraphs:
    for run in para.runs:
        if run.text.strip():
            run.text = "数据库构建漏斗 FUNNEL"
print("  P16 标题 ->", s16.shapes[2].text_frame.text.strip())
for para in s16.shapes[22].text_frame.paragraphs:
    for run in para.runs:
        if run.text.strip():
            run.text = "QED/SA/元素/环 395 0"
print("  P16 末级 ->", s16.shapes[22].text_frame.text.strip())

# ---------- P17 ----------
s17 = p.slides[16]
card = s17.shapes[17]
paras = card.text_frame.paragraphs
def set_para(para, txt):
    first = True
    for run in para.runs:
        if first:
            run.text = txt; first = False
        else:
            run.text = ""
set_para(paras[0], "403")
set_para(paras[1], "MW 中位 (Da)")
print("  P17 卡片4 ->", " / ".join(x.text for x in card.text_frame.paragraphs))

# ---------- P20 ----------
s20 = p.slides[19]
for para in s20.shapes[2].text_frame.paragraphs:
    for run in para.runs:
        if run.text.strip():
            run.text = "最大相似度中位 0.306、最大 0.549"
print("  P20 顶部 ->", s20.shapes[2].text_frame.text.strip())

# ---------- P22 框宽 ----------
s22 = p.slides[21]
for idx, w in [(31, 1.9), (27, 1.3), (42, 1.3), (39, 1.5)]:
    sh = s22.shapes[idx]
    old = Emu(sh.width).cm
    sh.width = Cm(w)
    print("  P22 [%d] 宽 %.2f -> %.2fcm" % (idx, old, w))

p.save(SRC)
print("SAVED", SRC)
