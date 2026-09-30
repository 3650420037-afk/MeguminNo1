# -*- coding: utf-8 -*-
"""用修好的尺寸重新生成的两张文字小图替换 PPT 中的图片。"""
import os, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.util import Emu

SRC = r'C:\Users\31908\Desktop\gpcr (1).pptx'
CH = r'D:\MMModel\outputs\night_pipeline_v10\charts'
p = Presentation(SRC)

def swap(slide, idx, png, tag):
    sh = slide.shapes[idx]
    left, top, w, h = sh.left, sh.top, sh.width, sh.height
    pic = slide.shapes.add_picture(os.path.join(CH, png), left, top, width=w, height=h)
    sh._element.addnext(pic._element)
    sh._element.getparent().remove(sh._element)
    print("%s: <- %s (%.2fx%.2fcm)" % (tag, png, Emu(w).cm, Emu(h).cm))

swap(p.slides[9], 13, "p10_chip.png", "P10 chip")
s22 = p.slides[21]
idx = [i for i, sh in enumerate(s22.shapes)
       if sh.shape_type is not None and 'PICTURE' in str(sh.shape_type)
       and abs(Emu(sh.top).cm - 9.49) < 0.1 and Emu(sh.left).cm < 3][0]
swap(s22, idx, "p22_chip.png", "P22 chip")
p.save(SRC)
print("SAVED")
