# -*- coding: utf-8 -*-
"""PPT 第二轮定点修正：行标签错位、残留旧数、B2AR 行清除、合计行补全。"""
import shutil, sys, os
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation

SRC = r'C:\Users\31908\Desktop\gpcr (1).pptx'
BAK = r'D:\MMModel\outputs\ppt_backup\gpcr_(1)_round2前_20260930.pptx'

def para_replace(tf, old, new, tag):
    """在 text_frame 内按段落做整段文本替换（保留首个 run 的格式）。"""
    hit = 0
    for para in tf.paragraphs:
        joined = "".join(r.text for r in para.runs)
        if old in joined:
            nb = joined.replace(old, new)
            if para.runs:
                para.runs[0].text = nb
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.text = nb
            hit += 1
    print("    %-28s %-34s -> %-34s hit=%d" % (tag, old, new, hit))
    return hit

def set_cell(cell, val, tag):
    tf = cell.text_frame
    if tf.paragraphs and tf.paragraphs[0].runs:
        tf.paragraphs[0].runs[0].text = val
        for r in tf.paragraphs[0].runs[1:]:
            r.text = ""
        for p in tf.paragraphs[1:]:
            for r in p.runs:
                r.text = ""
    else:
        tf.text = val
    print("    %-28s -> %r" % (tag, val))

if not os.path.exists(BAK):
    shutil.copy2(SRC, BAK)
    print("备份:", BAK)

p = Presentation(SRC)
S = p.slides

print("== P15 生成规模总览 ==")
tb = [sh for sh in S[14].shapes if sh.has_table][0].table
set_cell(tb.cell(1, 1), "D3 产出", "P15 t r1c1")
set_cell(tb.cell(1, 2), "5-HT2B 产出", "P15 t r1c2")
set_cell(tb.cell(1, 4), "", "P15 t r1c4")
set_cell(tb.cell(2, 4), "", "P15 t r2c4")
set_cell(tb.cell(3, 0), "22 批 × 50（实际 50-59/批）", "P15 t r3c0")
set_cell(tb.cell(5, 0), "D3 入库", "P15 t r5c0")
set_cell(tb.cell(6, 0), "5-HT2B 入库", "P15 t r6c0")
set_cell(tb.cell(7, 0), "三靶点合计", "P15 t r7c0")
set_cell(tb.cell(7, 2), "707", "P15 t r7c2")
para_replace(S[14].shapes[8].text_frame, "11 批", "5 批", "P15 [8]")

print("== P16 数据库漏斗 ==")
tb = [sh for sh in S[15].shapes if sh.has_table][0].table
set_cell(tb.cell(1, 2), "7", "P16 t r1c2")
set_cell(tb.cell(1, 3), "批间唯一率98.8%", "P16 t r1c3")

print("== P17 库质量画像 ==")
para_replace(S[16].shapes[20].text_frame, "56", "39", "P17 [20]")

print("== P20 新颖性验证 ==")
para_replace(S[19].shapes[2].text_frame, "最大相似度中位 0.242、最大 0.558 · 零骨架重合",
             "最大相似度中位 0.306、最大 0.549 · 骨架重合 1/395", "P20 [2]")
para_replace(S[19].shapes[2].text_frame, "最大相似度中位 0.306、最大 0.549 · 骨架重合 1/395",
             "最大相似度中位 0.306、最大 0.549 · 骨架重合 1/395", "P20 [2] 幂等")
para_replace(S[19].shapes[17].text_frame, "Tanimoto < 0.4 50", "Tanimoto < 0.4", "P20 [17]")
para_replace(S[19].shapes[18].text_frame, "100%", "94.7%", "P20 [18]")

print("== P22 跨靶点 ==")
for idx, old, new in [
    (31, "β2-AR", "D3 多巴胺"), (33, "-12.88", "-10.63"), (34, "38.9%", "5.7%"),
    (36, "D3多巴胺", "5-HT2B"), (37, "206 —", "206 -9.69"), (38, "-10.63", "-11.82"),
    (39, "5.7%", "37.4%"),
    (41, "5-HT2B —", "三靶点合计 707"), (42, "-9.69", "-9.70"), (43, "-11.82", "-12.88"),
    (44, "24.9%", "42.0%"),
]:
    para_replace(S[21].shapes[idx].text_frame, old, new, "P22 [%d]" % idx)
para_replace(S[21].shapes[45].text_frame, "四个 GPCR", "三个 GPCR", "P22 [45]")

print("== P25 数据卡 ==")
para_replace(S[24].shapes[30].text_frame, "<0.56", "<0.55", "P25 [30]")
para_replace(S[24].shapes[30].text_frame, "全部高新颖", "94% 分子 <0.4", "P25 [30]")

print("== P26 主张-证据 ==")
tb = [sh for sh in S[25].shapes if sh.has_table][0].table
set_cell(tb.cell(2, 1), "Top50 vs 737已知集 最大相似 0.495（均 <0.5）", "P26 t r2c1")
set_cell(tb.cell(4, 1), "三靶点1,151产出漏斗逐级可审", "P26 t r4c1")

print("== P30 瓶颈-方案-证据 ==")
tb = [sh for sh in S[29].shapes if sh.has_table][0].table
set_cell(tb.cell(2, 2), "Top50最大相似 0.495（均 <0.5）", "P30 t r2c2")
set_cell(tb.cell(3, 1), "三靶点数据管线+ 微调框架（接口 就绪）", "P30 t r3c1")
set_cell(tb.cell(3, 2), "三靶点1,151产出零失败", "P30 t r3c2")

print("== P31 附录总表 ==")
for idx, old, new in [
    (29, "β2-AR 254 106", "D3 254 106"), (31, "-12.88", "-10.63"), (32, "38.9% 49/50", "5.7% 50/50"),
    (34, "D3 314 206", "5-HT2B 314 206"), (35, "—", "-9.69"), (36, "-10.63", "-11.82"),
    (37, "5.7% 49/50", "37.4% 50/50"),
    (39, "5-HT2B — —", ""), (40, "-9.69", ""), (41, "-11.82", ""), (42, "24.9% 50/50", ""),
    (45, "、化合物库/", "、results/"), (46, "附录：四靶点成果总表", "附录：三靶点成果总表"),
    (48, "—", "-9.70"), (50, "— 150/150", "42.0% 150/150"),
]:
    para_replace(S[30].shapes[idx].text_frame, old, new, "P31 [%d]" % idx)

p.save(SRC)
print("\nSAVED", SRC)
