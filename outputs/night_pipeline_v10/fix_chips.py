# -*- coding: utf-8 -*-
"""替换两张文字小图与漏斗图；P15 卡片文字收尾。"""
import os, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.util import Emu
from PIL import Image

SRC = r'C:\Users\31908\Desktop\gpcr (1).pptx'
CH = r'D:\MMModel\outputs\night_pipeline_v10\charts'
p = Presentation(SRC)

def swap(slide, idx, png, tag):
    sh = slide.shapes[idx]
    left, top, w, h = sh.left, sh.top, sh.width, sh.height
    pic = slide.shapes.add_picture(os.path.join(CH, png), left, top, width=w, height=h)
    sh._element.addnext(pic._element)
    sh._element.getparent().remove(sh._element)
    print("%s: <- %s  (%.2fx%.2fcm @ %.2f,%.2f)" % (tag, png, Emu(w).cm, Emu(h).cm, Emu(left).cm, Emu(top).cm))

# P10 chip（“四靶点配对” -> “四口袋配对”）
swap(p.slides[9], 13, "p10_chip.png", "P10 chip")
# P22 chip（“同一管线 四靶点 零失败” -> “同一管线 · 三靶点 · 零失败”）
s22 = p.slides[21]
pic_idx = None
for i, sh in enumerate(s22.shapes):
    if sh.shape_type is not None and 'PICTURE' in str(sh.shape_type) and abs(Emu(sh.top).cm - 9.49) < 0.1 and Emu(sh.left).cm < 3:
        pic_idx = i
print("P22 chip 当前索引:", pic_idx)
swap(s22, pic_idx, "p22_chip.png", "P22 chip")

# P16 漏斗（第一次替换后的图片）
s16 = p.slides[15]
fidx = None
for i, sh in enumerate(s16.shapes):
    if sh.shape_type is not None and 'PICTURE' in str(sh.shape_type) and Emu(sh.height).cm > 8:
        fidx = i
print("P16 漏斗索引:", fidx)
swap(s16, fidx, "p16_funnel.png", "P16 funnel")

# P10 卡片文案
s10 = p.slides[9]
card = s10.shapes[21]
paras = card.text_frame.paragraphs
def set_para(para, txt):
    first = True
    for run in para.runs:
        if first:
            run.text = txt; first = False
        else:
            run.text = ""
set_para(paras[0], "12,244 对")
set_para(paras[1], "A2A · D3 · 5-HT2B")
for extra in paras[2:]:
    set_para(extra, "")
print("P10 卡片 ->", " | ".join(x.text for x in card.text_frame.paragraphs if x.text.strip()))

# P15 卡片文字
s15 = p.slides[14]
tb = [s for s in s15.shapes if s.has_table][0].table
def set_cell(tb, r, c, val):
    done = False
    for para in tb.cell(r, c).text_frame.paragraphs:
        for run in para.runs:
            if not done:
                run.text = val; done = True
            else:
                run.text = ""
    if not done:
        tb.cell(r, c).text_frame.text = val
set_cell(tb, 1, 4, "合计")
for idx, val in [(7, "5 批"), (8, "6 批"), (9, "22 批")]:
    sh = s15.shapes[idx]
    done = False
    for para in sh.text_frame.paragraphs:
        for run in para.runs:
            if not done:
                run.text = val; done = True
            else:
                run.text = ""
    print("P15 [%d] -> %s" % (idx, val))
print("P15 合计单元 ->", tb.cell(1, 4).text, "/", tb.cell(2, 4).text, "/", tb.cell(3, 0).text)

p.save(SRC)
print("SAVED", SRC)
