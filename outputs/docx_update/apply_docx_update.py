# -*- coding: utf-8 -*-
"""把《附件3.docx》区赛版二～六节替换为最新进展文本（保留原格式）。

- 只改 表1..表5（对应 二/三/四/五/六）的单元格段落；不动基本信息表、不动总决赛版。
- 段落级替换：复用原有段落对象以保留 pPr（对齐/缩进/字体继承），
  行数不足时深拷贝模板段落，多余段落删除。
- 打印每节字数（中文字符 + 英文/数字词）与模板上限，超限即报错退出（不写入）。
用法：
    python outputs/docx_update/apply_docx_update.py --dry-run
    python outputs/docx_update/apply_docx_update.py
"""
import argparse
import copy
import os
import re
import shutil
import sys
import time

import docx
from docx.text.paragraph import Paragraph

DOCX = r"C:\Users\31908\Desktop\附件3.docx"
HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "区赛二至六_新文本_20260927.txt")
LIMITS = {"二": 800, "三": 800, "四": 1800, "五": 3000, "六": 500}
TABLE_OF_SEC = {"二": 1, "三": 2, "四": 3, "五": 4, "六": 5}


def wc(text):
    zh = len(re.findall(r"[\u4e00-\u9fff]", text))
    en = len(re.findall(r"[A-Za-z0-9]+(?:[.\-/][A-Za-z0-9]+)*", text))
    return zh, en, zh + en


def load_sections(path):
    secs, cur, buf = {}, None, []
    for line in open(path, encoding="utf-8"):
        line = line.rstrip("\n")
        m = re.match(r"^@@SEC (\S+)$", line)
        if m:
            if cur:
                secs[cur] = buf
            cur, buf = m.group(1), []
            continue
        if cur is not None:
            buf.append(line)
    if cur:
        secs[cur] = buf
    return secs


def set_para_text(p, text):
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run(text)


def set_cell_lines(cell, lines):
    ps = cell.paragraphs
    if len(lines) <= len(ps):
        for i, ln in enumerate(lines):
            set_para_text(ps[i], ln)
        for i in range(len(ps) - 1, len(lines) - 1, -1):
            ps[i]._p.getparent().remove(ps[i]._p)
    else:
        for i, p in enumerate(ps):
            set_para_text(p, lines[i])
        tmpl = ps[-1]._p
        for i in range(len(ps), len(lines)):
            newp = copy.deepcopy(tmpl)
            tmpl.addnext(newp)
            tmpl = newp
            set_para_text(Paragraph(newp, cell), lines[i])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--docx", default=DOCX)
    args = ap.parse_args()

    secs = load_sections(SRC)
    print("读入节: %s" % sorted(secs))
    ok = True
    for k in ["二", "三", "四", "五", "六"]:
        lines = [l for l in secs[k]]
        text = "\n".join(lines)
        zh, en, tot = wc(text)
        lim = LIMITS[k]
        flag = "OK" if tot <= lim else "**超限**"
        ok = ok and tot <= lim
        print("  第%s节: 中文 %d + 英文/数字 %d = %d 字 (上限 %d) %s | 段落 %d"
              % (k, zh, en, tot, lim, flag, len(lines)))
    print("字数检查: %s" % ("全部通过" if ok else "存在超限，未写入"))
    if not ok:
        return 1

    d = docx.Document(args.docx)
    for k, ti in TABLE_OF_SEC.items():
        cell = d.tables[ti].cell(0, 0)
        old = len(cell.paragraphs)
        set_cell_lines(cell, secs[k])
        print("  表%d(第%s节): 段落 %d -> %d" % (ti, k, old, len(cell.paragraphs)))

    if args.dry_run:
        print("[dry-run] 未写入文件")
        return 0
    bak = os.path.join(HERE, "附件3_写入前_%s.docx" % time.strftime("%Y%m%d_%H%M%S"))
    shutil.copy2(args.docx, bak)
    d.save(args.docx)
    print("已写入 %s（写入前副本 %s）" % (args.docx, bak))
    return 0


if __name__ == "__main__":
    sys.exit(main())
